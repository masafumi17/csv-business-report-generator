"""
Excel chart builder module using openpyxl.
Generates line charts for daily sales trends and bar charts for rep/category sales.
"""

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.worksheet import Worksheet


def add_daily_sales_chart(ws: Worksheet, max_row: int, cell_target: str = "E2") -> None:
    """
    Add a line chart for daily sales trend into the specified worksheet.

    Args:
        ws (Worksheet): Target openpyxl worksheet (日別集計).
        max_row (int): Last row index of the data table.
        cell_target (str): Top-left cell coordinate to place the chart (e.g., 'E2').
    """
    if max_row < 2:
        return

    chart = LineChart()
    chart.title = "日別売上推移"
    chart.style = 13
    chart.y_axis.title = "売上金額 (円)"
    chart.x_axis.title = "注文日"
    chart.width = 16
    chart.height = 10

    # Data is in col 2 (売上金額)
    data = Reference(ws, min_col=2, min_row=1, max_row=max_row)
    # Categories are in col 1 (注文日)
    cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None  # Single series, no legend needed

    ws.add_chart(chart, cell_target)


def add_sales_rep_chart(ws: Worksheet, max_row: int, cell_target: str = "F2") -> None:
    """
    Add a bar chart for sales by representative into the worksheet.

    Args:
        ws (Worksheet): Target openpyxl worksheet (担当者別集計).
        max_row (int): Last row index of the data table.
        cell_target (str): Top-left cell coordinate to place the chart.
    """
    if max_row < 2:
        return

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "担当者別売上比較"
    chart.y_axis.title = "売上金額 (円)"
    chart.x_axis.title = "担当者"
    chart.width = 15
    chart.height = 10

    data = Reference(ws, min_col=2, min_row=1, max_row=max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None

    ws.add_chart(chart, cell_target)


def add_category_sales_chart(ws: Worksheet, max_row: int, cell_target: str = "F2") -> None:
    """
    Add a bar chart for sales by product category into the worksheet.

    Args:
        ws (Worksheet): Target openpyxl worksheet (カテゴリ別集計).
        max_row (int): Last row index of the data table.
        cell_target (str): Top-left cell coordinate to place the chart.
    """
    if max_row < 2:
        return

    chart = BarChart()
    chart.type = "col"
    chart.style = 11
    chart.title = "カテゴリ別売上比較"
    chart.y_axis.title = "売上金額 (円)"
    chart.x_axis.title = "商品カテゴリ"
    chart.width = 15
    chart.height = 10

    data = Reference(ws, min_col=2, min_row=1, max_row=max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None

    ws.add_chart(chart, cell_target)
